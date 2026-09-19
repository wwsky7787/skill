#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Build lightweight SQLite database with FTS5 trigram indexing from ICD Excel workbook.
Cross-platform compatible: Windows, Linux, macOS.
"""

import os
import sys
import sqlite3
from pathlib import Path
import openpyxl

def build_database(excel_path: Path, db_path: Path):
    if not excel_path.exists():
        raise FileNotFoundError(f"Excel file not found: {excel_path}")

    if db_path.exists():
        db_path.unlink()

    print(f"Loading Excel workbook: {excel_path.name}...")
    wb = openpyxl.load_workbook(str(excel_path), read_only=True, data_only=True)

    conn = sqlite3.connect(str(db_path))
    cur = conn.cursor()

    # Enable fast insertion mode
    cur.execute("PRAGMA journal_mode = OFF;")
    cur.execute("PRAGMA synchronous = 0;")
    cur.execute("PRAGMA cache_size = 100000;")

    # 1. Diseases
    print("Processing sheet: 湖北2.0版疾病完整库...")
    cur.execute("""
        CREATE TABLE diseases (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            code TEXT NOT NULL,
            name TEXT NOT NULL
        );
    """)
    cur.execute("CREATE INDEX idx_diseases_code ON diseases(code);")
    cur.execute("CREATE INDEX idx_diseases_name ON diseases(name);")

    ws_dis = wb["湖北2.0版疾病完整库"]
    dis_rows = []
    first = True
    for row in ws_dis.iter_rows(values_only=True):
        if first:
            first = False
            continue
        if not row or row[0] is None or row[2] is None:
            continue
        code = str(row[0]).strip()
        name = str(row[2]).strip()
        if code and name:
            dis_rows.append((code, name))

    cur.executemany("INSERT INTO diseases (code, name) VALUES (?, ?);", dis_rows)
    print(f"  Inserted {len(dis_rows)} disease records.")

    # FTS5 for Diseases
    cur.execute("""
        CREATE VIRTUAL TABLE diseases_fts USING fts5(
            name,
            content='diseases',
            content_rowid='id',
            tokenize=trigram
        );
    """)
    cur.execute("INSERT INTO diseases_fts(diseases_fts) VALUES('rebuild');")

    # 2. Surgeries
    print("Processing sheet: 湖北2.0版手术完整库...")
    cur.execute("""
        CREATE TABLE surgeries (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            code TEXT NOT NULL,
            name TEXT NOT NULL,
            level INTEGER
        );
    """)
    cur.execute("CREATE INDEX idx_surgeries_code ON surgeries(code);")
    cur.execute("CREATE INDEX idx_surgeries_name ON surgeries(name);")

    ws_surg = wb["湖北2.0版手术完整库"]
    surg_rows = []
    first = True
    for row in ws_surg.iter_rows(values_only=True):
        if first:
            first = False
            continue
        if not row or row[0] is None or row[2] is None:
            continue
        code = str(row[0]).strip()
        name = str(row[2]).strip()
        level_val = row[3]
        level = None
        if level_val is not None:
            try:
                level = int(level_val)
            except (ValueError, TypeError):
                level = None

        if code and name:
            surg_rows.append((code, name, level))

    cur.executemany("INSERT INTO surgeries (code, name, level) VALUES (?, ?, ?);", surg_rows)
    print(f"  Inserted {len(surg_rows)} surgery records.")

    # FTS5 for Surgeries
    cur.execute("""
        CREATE VIRTUAL TABLE surgeries_fts USING fts5(
            name,
            content='surgeries',
            content_rowid='id',
            tokenize=trigram
        );
    """)
    cur.execute("INSERT INTO surgeries_fts(surgeries_fts) VALUES('rebuild');")

    # Optimize & Vacuum
    print("Optimizing database & vacuuming...")
    conn.commit()
    cur.execute("INSERT INTO diseases_fts(diseases_fts) VALUES('optimize');")
    cur.execute("INSERT INTO surgeries_fts(surgeries_fts) VALUES('optimize');")
    cur.execute("PRAGMA optimize;")
    conn.commit()
    conn.close()

    # Reconnect to vacuum in normal mode
    conn = sqlite3.connect(str(db_path))
    conn.execute("VACUUM;")
    conn.close()

    db_size_mb = db_path.stat().st_size / (1024 * 1024)
    print(f"Database successfully created at: {db_path} ({db_size_mb:.2f} MB)")


if __name__ == "__main__":
    root_dir = Path(__file__).resolve().parent.parent
    if len(sys.argv) > 1:
        excel_file = Path(sys.argv[1]).resolve()
    else:
        # Search in root_dir first, then parent directory
        excel_candidates = list(root_dir.glob("*.xlsx")) + list(root_dir.parent.glob("*ICD*.xlsx"))
        if not excel_candidates:
            print("Error: No ICD Excel file found in", root_dir, "or", root_dir.parent)
            sys.exit(1)
        excel_file = excel_candidates[0]

    out_db = root_dir / "icd_hubei2.db"
    build_database(excel_file, out_db)
