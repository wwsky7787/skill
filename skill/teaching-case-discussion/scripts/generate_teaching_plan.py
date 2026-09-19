#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
教学病例讨论教案生成器
生成符合格式要求的 docx 文件教案

格式要求：
- A4 纸（210mm × 297mm）
- 版心尺寸：156mm × 225mm
- 页边距：上 36mm、下 36mm、左 27mm、右 27mm
- 每页 22 行，每行 28 字
- 正文：三号仿宋_GB2312，行间距 28 磅
- 标题层级格式按要求设置
"""

from docx import Document
from docx.shared import Pt, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn


def set_page_format(doc):
    """设置页面格式：A4 纸，版心尺寸，页边距"""
    section = doc.sections[0]
    
    # A4 纸尺寸
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    
    # 页边距：上 36mm、下 36mm、左 27mm、右 27mm
    section.top_margin = Mm(36)
    section.bottom_margin = Mm(36)
    section.left_margin = Mm(27)
    section.right_margin = Mm(27)
    
    return section


def set_font_style(run, font_name='仿宋_GB2312', font_size=16, bold=False, italic=False):
    """设置字体样式"""
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.italic = italic
    
    # 设置中文字体
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)


def set_paragraph_format(paragraph, line_spacing=28, space_before=0, space_after=0, alignment=WD_ALIGN_PARAGRAPH.LEFT):
    """设置段落格式"""
    p_format = paragraph.paragraph_format
    
    # 行间距（固定值 28 磅）
    p_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    p_format.line_spacing = Pt(line_spacing)
    
    # 段前段后间距
    p_format.space_before = Pt(space_before)
    p_format.space_after = Pt(space_after)
    
    # 对齐方式
    p_format.alignment = alignment
    
    return p_format


def add_section_title(doc, title):
    """添加章节标题（一级标题：一、二、三、）"""
    paragraph = doc.add_paragraph()
    set_paragraph_format(paragraph, line_spacing=28, space_before=24, space_after=12)
    
    run = paragraph.add_run(title)
    set_font_style(run, font_name='黑体', font_size=16, bold=True)
    
    return paragraph


def add_numbered_paragraph(doc, text, level=3, number=''):
    """添加带序号的段落"""
    paragraph = doc.add_paragraph()
    set_paragraph_format(paragraph, line_spacing=28, space_before=0, space_after=0)
    
    # 首行缩进
    paragraph.paragraph_format.first_line_indent = Pt(28 * 2)
    
    # 添加序号和内容
    if level == 1:
        # "一、"（三号黑体）
        run_num = paragraph.add_run(number)
        set_font_style(run_num, font_name='黑体', font_size=16, bold=True)
        run_text = paragraph.add_run(text)
        set_font_style(run_text, font_name='仿宋_GB2312', font_size=16)
        
    elif level == 2:
        # "(一)"（三号楷体）
        run_num = paragraph.add_run(number)
        set_font_style(run_num, font_name='楷体_GB2312', font_size=16)
        run_text = paragraph.add_run(text)
        set_font_style(run_text, font_name='仿宋_GB2312', font_size=16)
        
    elif level == 3:
        # "1."（三号仿宋，有标点）
        run_num = paragraph.add_run(number)
        set_font_style(run_num, font_name='仿宋_GB2312', font_size=16)
        run_text = paragraph.add_run(text)
        set_font_style(run_text, font_name='仿宋_GB2312', font_size=16)
        
    elif level == 4:
        # "(1)"（三号仿宋，必须有标点）
        run_num = paragraph.add_run(number)
        set_font_style(run_num, font_name='仿宋_GB2312', font_size=16)
        run_text = paragraph.add_run(text)
        set_font_style(run_text, font_name='仿宋_GB2312', font_size=16)
    
    return paragraph


def add_table(doc, data, column_widths=None):
    """添加表格"""
    rows = len(data)
    cols = len(data[0]) if data else 0
    
    table = doc.add_table(rows=rows, cols=cols)
    table.style = 'Table Grid'
    
    # 设置列宽
    if column_widths:
        for i, width in enumerate(column_widths):
            if i < cols:
                table.columns[i].width = Mm(width)
    
    # 填充数据
    for i, row_data in enumerate(data):
        row = table.rows[i]
        for j, cell_data in enumerate(row_data):
            cell = row.cells[j]
            cell.text = str(cell_data) if cell_data else ''
            
            # 设置单元格内文本格式
            for paragraph in cell.paragraphs:
                set_paragraph_format(paragraph, line_spacing=24)
                for run in paragraph.runs:
                    set_font_style(run, font_name='仿宋_GB2312', font_size=14)
    
    return table


def add_header_info(doc, info_dict):
    """添加基本信息表格"""
    # 标题
    title_para = doc.add_paragraph()
    set_paragraph_format(title_para, line_spacing=28, space_before=0, space_after=12, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    run = title_para.add_run('住院医师规范化培训教学病例讨论教案')
    set_font_style(run, font_name='黑体', font_size=22, bold=True)
    
    # 空一行
    doc.add_paragraph()
    
    # 基本信息表格
    info_data = [
        ['培训基地', info_dict.get('培训基地', '')],
        ['专业基地/科室', info_dict.get('专业基地/科室', '')],
        ['指导医师', info_dict.get('指导医师', '')],
        ['参加医师人数', info_dict.get('参加医师人数', '')],
        ['教学时长', info_dict.get('教学时长', '')],
        ['讨论日期', info_dict.get('讨论日期', '')],
    ]
    
    table = add_table(doc, info_data, column_widths=[40, 116])
    
    # 设置表格单元格对齐
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    return table


def generate_teaching_plan(jiao_an_data, output_path):
    """
    生成教案 docx 文件
    
    参数:
        jiao_an_data: dict，包含教案所有数据
        output_path: str，输出文件路径
    """
    doc = Document()
    
    # 设置页面格式
    set_page_format(doc)
    
    # 1. 添加标题和基本信息
    add_header_info(doc, jiao_an_data.get('基本信息', {}))
    
    # 2. 教学主题
    add_section_title(doc, '一、教学主题')
    zhu_ti_data = jiao_an_data.get('教学主题', {})
    add_numbered_paragraph(doc, '主题名称：' + zhu_ti_data.get('主题名称', ''), level=3, number='1.')
    add_numbered_paragraph(doc, '患者病历号：' + zhu_ti_data.get('患者病历号', ''), level=3, number='2.')
    add_numbered_paragraph(doc, '疾病名称：' + zhu_ti_data.get('疾病名称', ''), level=3, number='3.')
    
    # 3. 教学病例资料
    add_section_title(doc, '二、教学病例资料')
    bing_li_zi_liao = jiao_an_data.get('教学病例资料', {})
    
    add_numbered_paragraph(doc, '病例摘要：', level=3, number='1.')
    
    # 病例摘要内容
    add_numbered_paragraph(doc, '患者基本信息：' + bing_li_zi_liao.get('患者基本信息', ''), level=4, number='(1)')
    add_numbered_paragraph(doc, '主诉：' + bing_li_zi_liao.get('主诉', ''), level=4, number='(2)')
    add_numbered_paragraph(doc, '现病史：' + bing_li_zi_liao.get('现病史', ''), level=4, number='(3)')
    add_numbered_paragraph(doc, '既往史：' + bing_li_zi_liao.get('既往史', ''), level=4, number='(4)')
    add_numbered_paragraph(doc, '个人史：' + bing_li_zi_liao.get('个人史', ''), level=4, number='(5)')
    add_numbered_paragraph(doc, '家族史：' + bing_li_zi_liao.get('家族史', ''), level=4, number='(6)')
    add_numbered_paragraph(doc, '体格检查：' + bing_li_zi_liao.get('体格检查', ''), level=4, number='(7)')
    add_numbered_paragraph(doc, '辅助检查：' + bing_li_zi_liao.get('辅助检查', ''), level=4, number='(8)')
    add_numbered_paragraph(doc, '入院诊断：' + bing_li_zi_liao.get('入院诊断', ''), level=4, number='(9)')
    add_numbered_paragraph(doc, '治疗经过：' + bing_li_zi_liao.get('治疗经过', ''), level=4, number='(10)')
    
    # 讨论问题
    add_numbered_paragraph(doc, '讨论问题：', level=3, number='2.')
    tao_lun_wen_ti = jiao_an_data.get('讨论问题', [])
    for i, wen_ti in enumerate(tao_lun_wen_ti, 1):
        # 支持字典格式（包含问题和参考答案）或字符串格式
        if isinstance(wen_ti, dict):
            # 字典格式：{'问题': '...', '参考答案': '...'}
            add_numbered_paragraph(doc, wen_ti.get('问题', ''), level=4, number='(' + str(i) + ')')
            if wen_ti.get('参考答案'):
                add_numbered_paragraph(doc, '参考答案：' + wen_ti.get('参考答案', ''), level=4, number='(' + str(i) + ')')
        else:
            # 字符串格式：直接显示问题
            add_numbered_paragraph(doc, wen_ti, level=4, number='(' + str(i) + ')')
    
    # 4. 教学目标
    add_section_title(doc, '三、教学目标')
    jiao_xue_mu_biao = jiao_an_data.get('教学目标', [])
    for i, mu_biao in enumerate(jiao_xue_mu_biao, 1):
        add_numbered_paragraph(doc, mu_biao, level=3, number=str(i) + '.')
    
    # 重点知识和前沿知识
    if jiao_an_data.get('重点知识'):
        add_numbered_paragraph(doc, '重点知识：' + jiao_an_data.get('重点知识', ''), level=4, number='(' + str(len(jiao_xue_mu_biao)+1) + ')')
    if jiao_an_data.get('前沿知识'):
        add_numbered_paragraph(doc, '前沿知识：' + jiao_an_data.get('前沿知识', ''), level=4, number='(' + str(len(jiao_xue_mu_biao)+2) + ')')
    
    # 5. 课前准备
    add_section_title(doc, '四、课前准备')
    
    add_numbered_paragraph(doc, '指导医师准备：', level=3, number='1.')
    zhi_dao_jiao_shi_zhun_bei = jiao_an_data.get('指导医师准备', [])
    for item in zhi_dao_jiao_shi_zhun_bei:
        add_numbered_paragraph(doc, item, level=4, number='(1)')
    
    add_numbered_paragraph(doc, '住院医师准备：', level=3, number='2.')
    zhu_yuan_jiao_shi_zhun_bei = jiao_an_data.get('住院医师准备', [])
    for item in zhu_yuan_jiao_shi_zhun_bei:
        add_numbered_paragraph(doc, item, level=4, number='(1)')
    
    add_numbered_paragraph(doc, '场地与设备准备：', level=3, number='3.')
    chang_di_she_bei_zhun_bei = jiao_an_data.get('场地设备准备', [])
    for item in chang_di_she_bei_zhun_bei:
        add_numbered_paragraph(doc, item, level=4, number='(1)')
    
    # 6. 教学实施计划
    add_section_title(doc, '五、教学实施计划')
    
    shi_shi_ji_hua = jiao_an_data.get('教学实施计划', [])
    if shi_shi_ji_hua:
        # 表格数据
        table_data = [['时间', '环节', '内容', '教学形式', '重点/备注']]
        for item in shi_shi_ji_hua:
            table_data.append([
                item.get('时间', ''),
                item.get('环节', ''),
                item.get('内容', ''),
                item.get('教学形式', ''),
                item.get('重点/备注', '')
            ])
        add_table(doc, table_data, column_widths=[25, 30, 50, 25, 26])
    
    # 7. 评价与反馈计划
    add_section_title(doc, '六、评价与反馈计划')
    
    add_numbered_paragraph(doc, '住院医师评价：', level=3, number='1.')
    yi_shi_ping_jia = jiao_an_data.get('住院医师评价', [])
    for item in yi_shi_ping_jia:
        add_numbered_paragraph(doc, item, level=4, number='(1)')
    
    add_numbered_paragraph(doc, '课程评价：', level=3, number='2.')
    ke_cheng_ping_jia = jiao_an_data.get('课程评价', [])
    for item in ke_cheng_ping_jia:
        add_numbered_paragraph(doc, item, level=4, number='(1)')
    
    add_numbered_paragraph(doc, '课后作业：', level=3, number='3.')
    ke_hou_zuo_ye = jiao_an_data.get('课后作业', [])
    for item in ke_hou_zuo_ye:
        add_numbered_paragraph(doc, item, level=4, number='(1)')
    
    # 8. 参考资料
    add_section_title(doc, '七、参考资料')
    
    can_kao_zi_liao = jiao_an_data.get('参考资料', {})
    if can_kao_zi_liao.get('经典教材'):
        add_numbered_paragraph(doc, '经典教材：', level=3, number='1.')
        for item in can_kao_zi_liao.get('经典教材', []):
            add_numbered_paragraph(doc, item, level=4, number='(1)')
    
    if can_kao_zi_liao.get('临床指南'):
        add_numbered_paragraph(doc, '临床指南：', level=3, number='2.')
        for item in can_kao_zi_liao.get('临床指南', []):
            add_numbered_paragraph(doc, item, level=4, number='(1)')
    
    if can_kao_zi_liao.get('重要文献'):
        add_numbered_paragraph(doc, '重要文献：', level=3, number='3.')
        for item in can_kao_zi_liao.get('重要文献', []):
            add_numbered_paragraph(doc, item, level=4, number='(1)')
    
    # 保存文件
    doc.save(output_path)
    return output_path


if __name__ == '__main__':
    # 测试示例
    ce_shi_shu_ju = {
        '基本信息': {
            '培训基地': 'XX 医院',
            '专业基地/科室': '内科基地/呼吸内科',
            '指导医师': '张三 主任医师',
            '参加医师人数': '10 人',
            '教学时长': '60 分钟',
            '讨论日期': '2026 年 3 月 5 日'
        },
        '教学主题': {
            '主题名称': '肺炎链球菌肺炎的诊断与鉴别诊断',
            '患者病历号': '20260001',
            '疾病名称': '肺炎链球菌肺炎（Streptococcus pneumoniae pneumonia）'
        },
        '教学病例资料': {
            '患者基本信息': '男，65 岁，退休工人',
            '主诉': '发热、咳嗽 5 天，加重伴气促 2 天',
            '现病史': '患者 5 天前受凉后出现发热...',
            '既往史': '高血压病史 10 年',
            '个人史': '吸烟史 30 年',
            '家族史': '无特殊',
            '体格检查': 'T 39.2℃，右下肺可闻及湿啰音',
            '辅助检查': 'WBC 15.6×10^9/L，胸片示右下肺大片状阴影',
            '入院诊断': '社区获得性肺炎',
            '治疗经过': '给予抗感染、对症支持治疗'
        },
        '讨论问题': [
            '该患者的临床表现支持肺炎链球菌肺炎的诊断吗？',
            '需要与哪些疾病进行鉴别诊断？',
            '根据 CURB-65 评分，该患者的严重程度如何？'
        ],
        '教学目标': [
            '描述肺炎链球菌肺炎的典型临床表现',
            '运用临床思维方法进行诊断和鉴别诊断',
            '根据指南制定合理的抗感染治疗方案'
        ],
        '重点知识': '肺炎链球菌肺炎的诊断要点和鉴别诊断',
        '前沿知识': '2024 年 CAP 指南更新要点',
        '指导医师准备': ['选择病例', '明确目标', '查阅文献'],
        '住院医师准备': ['阅读资料', '查阅文献', '准备发言'],
        '场地设备准备': ['场地布置', '设备测试', '教具准备'],
        '教学实施计划': [
            {'时间': '5-10 分钟', '环节': '开场介绍', '内容': '介绍教学目标', '教学形式': '讲授式', '重点/备注': ''},
            {'时间': '30-35 分钟', '环节': '讨论与分析', '内容': '围绕问题展开', '教学形式': '互动式', '重点/备注': '核心环节'}
        ],
        '住院医师评价': ['病例资料掌握程度', '问题分析能力', '临床思维展现'],
        '课程评价': ['病例选择恰当性', '教学目标适合度', '指导医师引导能力'],
        '课后作业': ['撰写病例分析报告', '查阅指定文献'],
        '参考资料': {
            '经典教材': ['《内科学》（第 9 版）'],
            '临床指南': ['中华医学会社区获得性肺炎诊断和治疗指南（2024）'],
            '重要文献': ['肺炎链球菌肺炎诊治进展。中华结核和呼吸杂志，2024']
        }
    }
    
    generate_teaching_plan(ce_shi_shu_ju, '测试教案.docx')
    print('教案生成成功！')
